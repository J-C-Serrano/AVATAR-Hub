"""Rebuild Rodrigo Rabetino's 'Developing business models' slides as an editable .pptx from his recorded lecture.
usage (venv with python-pptx): python rod_deck.py <assets dir> <out.pptx> [--calib calib.json]
Geometry is in pixels of the 1920x1080 recording (measured from the frames, see HANDOVER), text sizes/spacing
were fitted to the recorded glyphs (Open Sans body, Georgia Bold titles). Shapes named 'u<k>_...' appear on the
k-th click (real PowerPoint 'Appear' animations, also used by deck_build.py for the video build-up).
--calib: per-style vertical offsets (px) measured from a test render, so ink tops land where the recording has them."""
import argparse, json, os
from copy import deepcopy
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ap = argparse.ArgumentParser(); ap.add_argument("assets"); ap.add_argument("out"); ap.add_argument("--calib")
ap.add_argument("--email", default="", help="contact line on the thank-you slide (kept out of the repo)")
a = ap.parse_args()
A = lambda f: os.path.join(a.assets, f)
CAL = json.load(open(a.calib)) if a.calib and os.path.exists(a.calib) else {}
PX = 6350  # EMU per pixel of the 1920-wide recording (13.333 in slide)
GRAY, BUL, TRI, RED, ELL, FOOT, NAME = "4D4D4D", "F2B533", "FCA400", "FF1A00", "CC1000", "898989", "FCB420"
OS, GE = "Open Sans", "Georgia"

prs = Presentation(); prs.slide_width, prs.slide_height = Emu(1920 * PX), Emu(1080 * PX)
BLANK = prs.slide_layouts[6]
E = lambda v: Emu(int(round(v * PX)))


def pic(s, name, f, x, y, w=None, h=None):
    """picture at its pixel size in the recording (python-pptx would assume 72 dpi otherwise)."""
    if w is None: w, h = Image.open(A(f)).size
    p = s.shapes.add_picture(A(f), E(x), E(y), E(w), E(h)); p.name = name; return p


def text(s, name, x, top, w, runs, size, font=OS, color=GRAY, bold=False, spc=0.0, line=None, align="l", lines=1, style=None):
    """runs: list of paragraphs, each a list of (text, {bold,color,underline}) or a str. top = ink top in px."""
    st = style or f"{font}/{size}"
    em = size * 2; off = CAL.get(st, 0.31 * em if font == OS else 0.21 * em)
    lh = (line or size * 1.25) * 2
    b = s.shapes.add_textbox(E(x), E(top - off), E(w), E(lh * lines + em)); b.name = name
    tf = b.text_frame; tf.word_wrap = True; tf.auto_size = None
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0; tf.vertical_anchor = MSO_ANCHOR.TOP
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER}[align]
        if line: p.line_spacing = Pt(line)
        for t, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run(); r.text = t; f = r.font
            f.name = font; f.size = Pt(size); f.bold = o.get("bold", bold); f.underline = o.get("underline", False)
            f.color.rgb = RGBColor.from_string(o.get("color", color))
            if spc: r._r.get_or_add_rPr().set("spc", str(int(round(spc * 100))))
    return b


def shape(s, name, kind, x, y, w, h, color, rot=0, line=None):
    sh = s.shapes.add_shape(kind, E(x), E(y), E(w), E(h)); sh.name = name; sh.rotation = rot
    if line: sh.fill.background(); sh.line.color.rgb = RGBColor.from_string(color); sh.line.width = Emu(int(line * PX))
    else: sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(color); sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def content_frame(s, title_lines, title_top):
    """template of the content slides: small logo top right, big logo bottom right, triangles, footer, title."""
    pic(s, "logo_top", "logo_tr.png", 1580, 47); pic(s, "logo_bottom", "logo_br.png", 1575, 855); pic(s, "triangles", "tri3.png", 106, 214)
    text(s, "footer", 400, 1008, 740, "DEVELOPING BUSINESS MODELS | RODRIGO RABETINO", 8.6, color=FOOT, align="c", style="foot")
    text(s, "title", 124, title_top, 1420, list(title_lines), 35, font=GE, bold=True,
         spc=2.5, line=39 if len(title_lines) > 1 else None, lines=len(title_lines), style="title")


def tri_bullet(s, name, x, y, big=False):
    w, h = (14, 16) if big else (10, 12)
    return shape(s, name, MSO_SHAPE.ISOSCELES_TRIANGLE, x, y, h, w, BUL, rot=90)


# 1 ─ title
s = prs.slides.add_slide(BLANK); pic(s, "background", "title_bg.png", 0, 0, 1920, 1080)
text(s, "title", 168, 611, 1358, ["Developing business", "models"], 44.5, color="F5AA00", spc=4.6, line=60, align="c", lines=2, style="title1")
# 2 ─ section title
s = prs.slides.add_slide(BLANK); pic(s, "logo_top", "logo_tr.png", 1580, 25)
text(s, "title", 160, 239, 1365, ["Business", "models &", "business model", "innovation"], 73.5, font=GE, bold=True, spc=6.7, line=78.6,
     align="c", lines=4, style="title2")
pic(s, "triangles", "tri9.png", 674, 876)
text(s, "name", 500, 933, 690, "RODRIGO RABETINO", 16.5, color=NAME, bold=True, spc=4.1, align="c", style="name")
# 3 ─ learning outcomes (no build)
s = prs.slides.add_slide(BLANK); content_frame(s, ["Content & Intended Learning", "Outcomes"], 55)
tri_bullet(s, "bullet", 108, 316, big=True)
text(s, "intro", 140, 295, 1200, "By the end of the class, you will be able to understand:", 28, line=42.5, lines=2, style="ilo")
for k, (t, top) in enumerate([("What is value?", 543), ("What is a business model?", 645), ("The difference between Strategy and Business Models.", 748)]):
    text(s, f"num{k + 1}", 104, top + 11, 70, ["I.", "II.", "III."][k], 21.5, color=BUL, style="ilonum")
    text(s, f"item{k + 1}", 181, top, 1240, t, 28, line=42, lines=2, style="ilo")


def concept(s, built):
    """'Business models: The concept'; built = paragraphs visible from the start (rest appear on clicks)."""
    content_frame(s, ["Business models: The concept"], 133)
    P = [("lvl0", "It is important in any market to get the business model right.", 288),
         ("lvl0", "The role of innovation.", 403),
         ("lvl1", "Think about Kodak, Nokia, Agfa, and Grundig. What do they have in common?", 487),
         ("lvl1", "They failed to adapt their business models to the changing environment,", 641),
         ("lvl2", "Future competition will take place between business models, and not just between products and technologies,", 733),
         ("lvl2", "Business model innovation (BMI) to be more important for achieving competitive advantage than product or service innovation.", 875)]
    k = 0
    for i, (lv, t, top) in enumerate(P):
        pre = "" if i < built else f"u{(k := k + 1)}_"
        if lv == "lvl0": tri_bullet(s, pre + f"bullet{i}", 106, top + 16); text(s, pre + f"p{i}", 139, top, 1290, t, 20.8, style="lvl0")
        if lv == "lvl1": pic(s, pre + f"bullet{i}", "check.png", 172, top - 1); text(s, pre + f"p{i}", 228, top, 1195, t, 17.5, line=24.5, lines=2, style="lvl1")
        if lv == "lvl2": shape(s, pre + f"bullet{i}", MSO_SHAPE.RECTANGLE, 248, top + 9, 8, 8, BUL); text(s, pre + f"p{i}", 283, top, 1140, t, 15.8, line=24, lines=2, style="lvl2")


# 4 ─ concept, part 1 (appears with 2 bullets, Kodak on click) · 5 ─ clock · 6 ─ concept, part 2
s = prs.slides.add_slide(BLANK); concept(s, 2)
for sh in list(s.shapes):   # part 1 ends at the Kodak question
    if sh.name.split("_")[0] in ("u2", "u3", "u4"): sh._element.getparent().remove(sh._element)
s = prs.slides.add_slide(BLANK)
mv = s.shapes.add_movie(A("clock.mp4"), 0, 0, prs.slide_width, prs.slide_height, poster_frame_image=A("clock_poster.png"), mime_type="video/mp4"); mv.name = "clock"
s = prs.slides.add_slide(BLANK); concept(s, 3)
# 7 ─ trendsetters
s = prs.slides.add_slide(BLANK); content_frame(s, ["Trendsetters signals"], 133)
pic(s, "bullet0", "arrow.png", 100, 291)
R = {"bold": True, "color": RED, "underline": True}
text(s, "p0", 139, 291, 940, [[("New business models are often based on ", {}), ("early", R), (" weak ", R), ("signals:", {})]], 20.8, line=31.5, lines=2, style="lvl0t")
for k, (t, top) in enumerate([("New customer requirements;", 515), ("Regulations are discussed broadly before they are eventually  approved;", 655),
                              ("Disruptive technology developments are result of many years of research;", 850)]):
    pic(s, f"u{k + 1}_bullet", "check.png", 172, top - 1); text(s, f"u{k + 1}_p", 230, top, 830, t, 17.5, line=26.5, lines=2, style="lvl1t")
pic(s, "u4_book", "mcgrath.png", 1128, 335)
# 8 ─ importance of strategy
s = prs.slides.add_slide(BLANK); content_frame(s, ["The importance of strategy and", "business models"], 55)
pic(s, "u1_photos", "kaplan_norton.png", 1030, 268); pic(s, "u2_books", "books.png", 740, 590)
shape(s, "u3_circle", MSO_SHAPE.OVAL, 725, 655, 155, 307, ELL, line=5)
tri_bullet(s, "u4_bullet", 106, 304); text(s, "u4_p", 139, 288, 700, "85% of executive teams spend less than 1 hour / month to discuss strategy.", 20.8, line=28, lines=3, style="lvl0s")
tri_bullet(s, "u5_bullet", 106, 534); text(s, "u5_p", 139, 518, 700, "Only 5% of employees understand the strategy.", 20.8, line=28, lines=2, style="lvl0s")
# 9 ─ thank you
s = prs.slides.add_slide(BLANK); pic(s, "logo_top", "logo_ty.png", 1664, 34)
text(s, "name", 80, 128, 260, ["RODRIGO", "RABETINO"], 16.5, color=NAME, bold=True, spc=2.0, line=24, align="c", lines=2, style="name9")
pic(s, "triangles", "tri3_ty_left.png", 108, 216)
text(s, "thanks", 460, 342, 980, "THANK YOU!!!", 33, bold=True, spc=6.7, align="c", style="thanks")
text(s, "email", 460, 446, 1000, a.email.upper(), 18, color=NAME, bold=True, spc=1.5, align="c", style="email")
pic(s, "triangles2", "tri3_ty.png", 895, 504); pic(s, "people", "people.png", 60, 668)

# ── real 'Appear' animations from the u<k>_ names ──
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
for s in prs.slides:
    units = {}
    for sh in s.shapes:
        if sh.name[:1] == "u" and "_" in sh.name and sh.name.split("_")[0][1:].isdigit():
            units.setdefault(int(sh.name.split("_")[0][1:]), []).append(sh.shape_id)
    if not units: continue
    ids = iter(range(3, 10000))
    clicks = ""
    for k in sorted(units):
        effs = ""
        for j, spid in enumerate(units[k]):
            effs += (f'<p:par><p:cTn id="{next(ids)}" presetID="1" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" '
                     f'nodeType="{"clickEffect" if j == 0 else "withEffect"}"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
                     f'<p:set><p:cBhvr><p:cTn id="{next(ids)}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
                     f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>'
                     f'</p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set></p:childTnLst></p:cTn></p:par>')
        clicks += (f'<p:par><p:cTn id="{next(ids)}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
                   f'<p:par><p:cTn id="{next(ids)}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{effs}'
                   f'</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')
    bld = "".join(f'<p:bldP spid="{i}" grpId="0"/>' for k in units for i in units[k]
                  if s.shapes._spTree.xpath(f'./p:sp[p:nvSpPr/p:cNvPr/@id="{i}"]/p:txBody'))
    xml = (f'<p:timing xmlns:p="{P_NS}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>{clicks}</p:childTnLst></p:cTn>'
           f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
           f'</p:childTnLst></p:cTn></p:par></p:tnLst>{"<p:bldLst>" + bld + "</p:bldLst>" if bld else ""}</p:timing>')
    el = etree.fromstring(xml); sld = s._element
    anchor = sld.find(qn("p:clrMapOvr"))
    (anchor.addnext(el) if anchor is not None else sld.find(qn("p:cSld")).addnext(el))
prs.save(a.out); print("WROTE", a.out, len(prs.slides), "slides")
